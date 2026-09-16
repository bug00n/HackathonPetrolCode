from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


FILES = [
    Path(r"D:\downloads\Telegram Desktop\теги АВТ_24-2000.xlsx"),
    Path(r"D:\downloads\Telegram Desktop\формулы_ВАК.xlsx"),
]
TERMS = ()


for path in FILES:
    print(f"### {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    print("sheets", workbook.sheetnames)
    for sheet in workbook.worksheets:
        print(f"## {sheet.title}: rows={sheet.max_row}, cols={sheet.max_column}")
        rows = list(sheet.iter_rows(values_only=True))
        for index, row in enumerate(rows[:8], start=1):
            print("HEAD", index, json.dumps(row, ensure_ascii=False, default=str))
        for index, row in enumerate(rows, start=1):
            text = " | ".join("" if value is None else str(value) for value in row).lower()
            if not TERMS or any(term in text for term in TERMS):
                print("MATCH", index, json.dumps(row, ensure_ascii=False, default=str))
    workbook.close()
