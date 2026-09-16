from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path.cwd()))
from source.ml.formulas import apply_expert_vak_corrections, load_vak_formulas


def normalize(value: str) -> str:
    value = value.replace("−", "-").replace("×", "*").replace(" ", "")
    return re.sub(r"LIMS\.95%\.T", "LIMS:24-2000.Pipeline.95%.T", value)


old_path = Path("materials/Теги_хакатон.xlsx")
new_path = Path(r"D:\downloads\Telegram Desktop\формулы_ВАК.xlsx")
old = apply_expert_vak_corrections(load_vak_formulas(old_path))
new_book = load_workbook(new_path, read_only=True, data_only=False)
new: dict[str, str] = {}
for sheet in new_book.worksheets:
    for row in sheet.iter_rows(min_row=2, values_only=True):
        _, model, formula, *_ = row
        if model and formula:
            new[str(model)] = str(formula)

for model, formula in new.items():
    old_formula = old.get(model)
    status = (
        "MATCH"
        if old_formula is not None and normalize(old_formula) == normalize(formula)
        else "DIFF"
    )
    print(status, model)
    if status == "DIFF":
        print("  OLD", old_formula)
        print("  NEW", formula)
print("new models", len(new), "old models", len(old), "missing in old", sorted(set(new) - set(old)))
